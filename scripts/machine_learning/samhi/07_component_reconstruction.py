"""Build and audit the three non-hospital SAMHI components.

The hospital component in SAMHI v5 is safeguarded NHS data and is therefore
left explicitly missing here.  This script normalises the newly supplied
PLDR/DWP/QOF assets into an LSOA11-year panel so that the three observable
components can be checked, standardised, and later combined with an approved
hospital extract or an evaluated EMAS proxy.
"""

from __future__ import annotations

import argparse
import csv
import json
import importlib.util
import logging
import re
from pathlib import Path
from typing import Dict, List

import numpy as np
import pandas as pd

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
LOGGER = logging.getLogger("SAMHI_Component_Reconstruction")


def project_root() -> Path:
    return Path(__file__).resolve().parents[3]


def normalise_code(value: object) -> str:
    return str(value).split(":", 1)[0].strip()


def load_antidepressants(source_root: Path) -> pd.DataFrame:
    """Aggregate quarterly PLDR ADQ rates to annual LSOA observations.

    ``adqs_r`` is the PLDR rate per 1,000 population.  The SAMHI component is
    ADQ per person, so the annual value is the mean of quarterly ``adqs_r``
    divided by 1,000.  ``adqs / all_ages`` is retained as a diagnostic check.
    """
    base = source_root / "pldr_prescribing_indicators_antidepressants_P_1_07"
    rows: List[pd.DataFrame] = []
    for year_dir in sorted(base.iterdir() if base.exists() else []):
        if not year_dir.is_dir() or not re.fullmatch(r"\d{4}", year_dir.name):
            continue
        year = int(year_dir.name)
        for path in sorted(year_dir.glob("*_LSOA.csv")):
            raw = pd.read_csv(path)
            required = {"lsoa11", "adqs_r", "adqs", "all_ages"}
            if not required.issubset(raw.columns):
                LOGGER.warning("Skipping %s; missing columns %s", path.name, required - set(raw.columns))
                continue
            out = raw[["lsoa11", "adqs_r", "adqs", "all_ages"]].copy()
            out["lsoa11"] = out["lsoa11"].map(normalise_code)
            out = out[out["lsoa11"].str.startswith("E", na=False)].copy()
            for col in ["adqs_r", "adqs", "all_ages"]:
                out[col] = pd.to_numeric(out[col], errors="coerce")
            out["antidep_rate"] = out["adqs_r"] / 1000.0
            out["antidep_rate_check"] = out["adqs"] / out["all_ages"].replace(0, np.nan)
            out["year"] = year
            rows.append(out[["lsoa11", "year", "antidep_rate", "antidep_rate_check"]])
    if not rows:
        return pd.DataFrame(columns=["lsoa11", "year", "antidep_rate", "antidep_rate_check", "antidep_quarters"])
    out = pd.concat(rows, ignore_index=True)
    result = out.groupby(["lsoa11", "year"], as_index=False).agg(
        antidep_rate=("antidep_rate", "mean"),
        antidep_rate_check=("antidep_rate_check", "mean"),
        antidep_quarters=("antidep_rate", "count"),
    )
    LOGGER.info("Loaded antidepressants: %s LSOA-year rows, years %s-%s", len(result), result.year.min(), result.year.max())
    return result


def load_qof_pldr(source_root: Path) -> pd.DataFrame:
    """Load PLDR's corrected LSOA QOF-depression estimates for 2011--2022."""
    path = source_root / "pldr_qof_indicators_depression_prevalence_QOF_4_12" / "QOF_4_12_Depression_LSOA_2011_2022.csv"
    if not path.exists():
        return pd.DataFrame(columns=["lsoa11", "year", "qof_dep_pct"])
    raw = pd.read_csv(path)
    raw["lsoa11"] = raw["lsoa11"].map(normalise_code)
    raw["year"] = pd.to_numeric(raw["year"], errors="coerce").astype("Int64")
    # est_qof_dep contains PLDR's correction for the historical definition
    # break; retain the uncorrected value for sensitivity analysis.
    out = raw.rename(columns={"est_qof_dep": "qof_dep_pct", "qof_dep": "qof_dep_raw_pct"})
    out = out[out["lsoa11"].str.startswith("E", na=False)].copy()
    cols = ["lsoa11", "year", "qof_dep_pct", "qof_dep_raw_pct", "den", "num", "dum"]
    cols = [col for col in cols if col in out.columns]
    out = out[cols]
    LOGGER.info("Loaded PLDR QOF depression: %s LSOA-year rows, years %s-%s", len(out), out.year.min(), out.year.max())
    return out


def _read_qof_depression_practice(path: Path, year: int) -> pd.DataFrame:
    """Read practice depression prevalence from a public QOF workbook.

    Recent QOF workbooks contain a ``DEP`` sheet.  The prevalence columns are
    not stable: some years publish two columns (previous/current year), while
    2023--24 publishes no depression prevalence at all.  We therefore select
    the final prevalence column when it exists and return an empty frame when
    the workbook only contains achievement/incidence measures.
    """
    import openpyxl

    columns = ["practice_code", "qof_dep_practice_pct", "qof_year"]
    if not path.exists():
        return pd.DataFrame(columns=columns)
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "DEP" not in book.sheetnames:
        book.close()
        return pd.DataFrame(columns=columns)
    sheet = book["DEP"]
    header_row = None
    header = None
    for row_number, raw in enumerate(sheet.iter_rows(values_only=True), 1):
        values = [str(value).strip() if value is not None else "" for value in raw]
        lowered = {value.lower() for value in values}
        if "practice code" in lowered and any(value.lower().startswith("prevalence") for value in values):
            header_row, header = row_number, values
            break
    if header_row is None or header is None:
        book.close()
        return pd.DataFrame(columns=columns)
    practice_idx = next(idx for idx, value in enumerate(header) if value.lower() == "practice code")
    prevalence_indices = [idx for idx, value in enumerate(header) if value.lower().startswith("prevalence")]
    prevalence_idx = prevalence_indices[-1]
    records = []
    for raw in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if practice_idx >= len(raw):
            continue
        practice = normalise_code(raw[practice_idx])
        if not practice or practice.lower() in {"none", "nan", "practice code"}:
            continue
        value = raw[prevalence_idx] if prevalence_idx < len(raw) else np.nan
        prevalence = pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]
        if pd.notna(prevalence):
            records.append((practice, float(prevalence), year))
    book.close()
    return pd.DataFrame(records, columns=columns).drop_duplicates("practice_code", keep="last")


def load_qof_practice_lsoa(root: Path) -> pd.DataFrame:
    """Allocate public practice QOF depression prevalence to LSOA11.

    The allocation is patient-weighted using the repository's public GP
    practice-to-LSOA registered-patient extracts.  The output is deliberately
    separate from PLDR's historical LSOA estimate so that the two definitions
    can be compared before the public series is used for forecasting.
    """
    source_root = root / "scripts" / "utils" / "source_data"
    qof_base = source_root / "quality_outcomes_framework"
    datasets = root / "datasets" / "patients_registered_gp_practice" / "england"
    workbook_paths = {
        2021: qof_base / "2021-22" / "mental_health_neurology_group" / "qof-2122-prev-ach-pca-neu-prac.xlsx",
        2022: qof_base / "2022-23" / "mental_health_neurology_group" / "qof-2223-prev-ach-pca-neu-prac.xlsx",
        2023: qof_base / "2023-24" / "mental_health_neurology_group" / "qof-2324-prev-ach-pca-neu-prac.xlsx",
        2024: qof_base / "2024-25" / "mental_health_neurology_group" / "qof-2425-prev-ach-pca-neu-prac.xlsx",
    }
    output = []
    for year, workbook in workbook_paths.items():
        practice = _read_qof_depression_practice(workbook, year)
        if practice.empty:
            LOGGER.warning("No public QOF depression prevalence found for %s (%s)", year, workbook.name)
            continue
        patient_path = datasets / str(year + 1) / "july" / "gp-reg-pat-prac-lsoa-all.csv"
        if not patient_path.exists():
            LOGGER.warning("No July GP patient allocation found for QOF %s: %s", year, patient_path)
            continue
        patients = pd.read_csv(
            patient_path,
            usecols=["PRACTICE_CODE", "LSOA_CODE", "SEX", "NUMBER_OF_PATIENTS"],
            low_memory=False,
        )
        patients = patients[patients["SEX"].astype(str).str.upper().eq("ALL")].copy()
        patients["practice_code"] = patients["PRACTICE_CODE"].map(normalise_code)
        patients["lsoa11"] = patients["LSOA_CODE"].map(normalise_code)
        patients["patients"] = pd.to_numeric(patients["NUMBER_OF_PATIENTS"], errors="coerce")
        patients = patients[patients["lsoa11"].str.startswith("E", na=False) & patients["patients"].gt(0)]
        patients = patients.merge(practice, on="practice_code", how="inner")
        if patients.empty:
            LOGGER.warning("QOF %s had no practice-to-LSOA matches", year)
            continue
        patients["weighted_prevalence"] = patients["patients"] * patients["qof_dep_practice_pct"]
        lsoa = patients.groupby("lsoa11", as_index=False).agg(
            qof_public_pct=("weighted_prevalence", "sum"),
            qof_patient_weight=("patients", "sum"),
            qof_practice_count=("practice_code", "nunique"),
        )
        lsoa["qof_public_pct"] = lsoa["qof_public_pct"] / lsoa["qof_patient_weight"]
        # QOF labels the financial year by its start year (e.g. 2024-25),
        # while the component panel uses the year ending in March.  Preserve
        # both: ``year`` is the panel year and ``qof_public_source_year`` is
        # the workbook's start-year vintage.
        lsoa["year"] = year + 1
        lsoa["qof_public_source_year"] = year
        output.append(lsoa[["lsoa11", "year", "qof_public_source_year", "qof_public_pct", "qof_patient_weight", "qof_practice_count"]])
        LOGGER.info("Reconstructed public QOF %s: %s LSOAs from %s practices", year, len(lsoa), len(practice))
    if not output:
        return pd.DataFrame(columns=["lsoa11", "year", "qof_public_source_year", "qof_public_pct", "qof_patient_weight", "qof_practice_count"])
    return pd.concat(output, ignore_index=True)


def load_welfare_pldr(source_root: Path) -> pd.DataFrame:
    """Load annual August PLDR DLA/PIP claimant counts and normalise year."""
    path = source_root / "pldr_welfare_indicators_claimants_DLA_PIP_for_mental_health_learning_difficulties_W_5_05" / "W_5_05_MH_DLA_PIP_LSOA11.csv"
    if not path.exists():
        return pd.DataFrame(columns=["lsoa11", "year", "dla_pip"])
    raw = pd.read_csv(path)
    out = raw.rename(columns={"LSOA11CD": "lsoa11"})
    out["lsoa11"] = out["lsoa11"].map(normalise_code)
    out["year"] = (pd.to_numeric(out["year"], errors="coerce") // 100).astype("Int64")
    out["dla_pip"] = pd.to_numeric(out["dla_pip"], errors="coerce")
    out = out[out["lsoa11"].str.startswith("E", na=False)].copy()
    out = out[["lsoa11", "year", "dla_pip", "dla_pop", "pip_pop"]]
    LOGGER.info("Loaded PLDR DLA/PIP: %s LSOA-year rows, years %s-%s", len(out), out.year.min(), out.year.max())
    return out


def _load_population_helper(root: Path):
    """Load the existing population parser without duplicating its workbook logic."""
    path = root / "scripts" / "machine_learning" / "samhi" / "06_historical_bayesian_experiment.py"
    spec = importlib.util.spec_from_file_location("historical_experiment", path)
    if spec is None or spec.loader is None:
        raise ImportError(f"Could not load {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _read_dwp_statxplore_august(path: Path, name_to_lsoa21: Dict[str, str]) -> pd.DataFrame:
    """Read an August snapshot from a wide Stat-Xplore CSV export."""
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    header_idx = next((idx for idx, row in enumerate(rows) if row and row[0] in {"Month", "Quarter"}), None)
    if header_idx is None:
        return pd.DataFrame(columns=["lsoa21", "year", "value"])
    header = rows[header_idx]
    period_columns = []
    for idx, label in enumerate(header[1:], start=1):
        label = str(label).strip()
        if label and "annotation" not in label.lower() and label.startswith("Aug-"):
            period_columns.append((idx, label))
    output = []
    for row in rows[header_idx + 1:]:
        if not row or row[0] not in name_to_lsoa21:
            continue
        for idx, label in period_columns:
            value = row[idx] if idx < len(row) else ""
            numeric = pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]
            year = int(pd.to_datetime(label, format="%b-%y").year)
            output.append((name_to_lsoa21[row[0]], year, numeric))
    return pd.DataFrame(output, columns=["lsoa21", "year", "value"])


def load_dwp_welfare(source_root: Path, root: Path) -> pd.DataFrame:
    """Combine DWP DLA and PIP August counts after the PLDR series ends.

    The DWP files are 2021-LSOA, wide Stat-Xplore extracts.  Counts are joined
    to the national 2021 boundary names and mapped to LSOA11.  An optional
    denominator-normalised rate uses the estimated 16--64 population.  A 2024
    population denominator is used for the 2025 August snapshot because
    mid-2025 population estimates are not in the repository; that vintage is
    retained explicitly.
    """
    dwp_dir = source_root / "DWP_DLA_PIP_data"
    dla_path = dwp_dir / "DLA_claimants_mental_health_learning_difficulties_from_may_2018_lsoa.csv"
    pip_path = dwp_dir / "PIP_cases_with_entitlement_mental_health_learning_difficulties_from_2019_lsoa.csv"
    boundary_path = root / "datasets" / "england_lsoa" / "Lower_layer_Super_Output_Areas_December_2021_Boundaries_EW_BSC_V4_6894679968818356315.geojson"
    if not dla_path.exists() or not pip_path.exists() or not boundary_path.exists():
        return pd.DataFrame(columns=["lsoa11", "year", "dwp_dla_pip_count", "dwp_dla_pip_rate_pct", "dwp_dla_count", "dwp_pip_count", "dwp_population_source_year"])
    with boundary_path.open() as handle:
        geojson = json.load(handle)
    name_to_lsoa21 = {
        str(feature["properties"]["LSOA21NM"]): str(feature["properties"]["LSOA21CD"])
        for feature in geojson["features"]
        if feature.get("properties", {}).get("LSOA21NM")
    }
    dla = _read_dwp_statxplore_august(dla_path, name_to_lsoa21).rename(columns={"value": "dwp_dla_count"})
    pip = _read_dwp_statxplore_august(pip_path, name_to_lsoa21).rename(columns={"value": "dwp_pip_count"})
    if dla.empty or pip.empty:
        return pd.DataFrame(columns=["lsoa11", "year", "dwp_dla_pip_count", "dwp_dla_pip_rate_pct", "dwp_dla_count", "dwp_pip_count", "dwp_population_source_year"])
    out = dla.merge(pip, on=["lsoa21", "year"], how="inner")
    out["dwp_dla_pip_count"] = out[["dwp_dla_count", "dwp_pip_count"]].sum(axis=1, min_count=2)
    lookup_path = root / "datasets" / "lincolnshire_lsoa" / "lsoa_2011_to_2021_lookup" / "LSOA_(2011)_to_LSOA_(2021)_to_Local_Authority_District_(2022)_Exact_Fit_Lookup_for_EW_(V3).csv"
    # Use the national lookup copy if available; otherwise retain 2021 codes
    # and let the caller restrict to geographies with a valid crosswalk.
    if lookup_path.exists():
        lookup = pd.read_csv(lookup_path, usecols=["LSOA11CD", "LSOA21CD"])
        lookup = lookup.rename(columns={"LSOA11CD": "lsoa11", "LSOA21CD": "lsoa21"})
        lookup["lsoa11"] = lookup["lsoa11"].map(normalise_code)
        lookup["lsoa21"] = lookup["lsoa21"].map(normalise_code)
        out = out.merge(lookup.drop_duplicates("lsoa21"), on="lsoa21", how="inner")
    else:
        out["lsoa11"] = out["lsoa21"]
    # DWP snapshots begin in 2019; limiting the parser to the required
    # denominator vintages keeps this audit reproducible without rereading the
    # legacy 2011--2018 workbooks on every run.
    population = _load_population_helper(root).load_historical_population(source_root, min_year=2019, max_year=2024)
    population = population.rename(columns={"source_year": "population_source_year"})
    population = population[["lsoa21", "population_source_year", "pop_total", "pop_working_age_pct"]]
    # Select the latest available denominator not later than each DWP year.
    out = out.merge(population, on="lsoa21", how="left")
    out = out[out["population_source_year"] <= out["year"]].copy()
    out = out.sort_values(["lsoa21", "year", "population_source_year"]).drop_duplicates(["lsoa21", "year"], keep="last")
    denominator = out["pop_total"] * out["pop_working_age_pct"] / 100.0
    out["dwp_dla_pip_rate_pct"] = out["dwp_dla_pip_count"] / denominator.replace(0, np.nan) * 100.0
    out["dwp_population_source_year"] = out["population_source_year"]
    return out[["lsoa11", "year", "dwp_dla_pip_count", "dwp_dla_pip_rate_pct", "dwp_dla_count", "dwp_pip_count", "dwp_population_source_year"]]


def zscore_by_year(panel: pd.DataFrame, column: str) -> pd.Series:
    return panel.groupby("year")[column].transform(lambda values: (values - values.mean()) / values.std(ddof=0))


def load_published_samhi(datasets: Path) -> pd.DataFrame:
    path = datasets / "samhi" / "samhi_21_01_v5.00_2011_2022_LSOA.csv"
    raw = pd.read_csv(path)
    raw["lsoa11"] = raw["lsoa11"].map(normalise_code)
    rows = []
    for year in range(2011, 2023):
        col = f"samhi_index.{year}"
        if col in raw.columns:
            rows.append(raw[["lsoa11", col]].rename(columns={col: "samhi_index"}).assign(year=year))
    return pd.concat(rows, ignore_index=True) if rows else pd.DataFrame(columns=["lsoa11", "year", "samhi_index"])


def build_panel(root: Path) -> pd.DataFrame:
    source_root = root / "scripts" / "utils" / "source_data"
    components = [load_antidepressants(source_root), load_qof_pldr(source_root), load_welfare_pldr(source_root)]
    panel = components[0]
    for component in components[1:]:
        panel = panel.merge(component, on=["lsoa11", "year"], how="outer")
    panel = panel[panel["lsoa11"].str.startswith("E", na=False)].copy()
    dwp = load_dwp_welfare(source_root, root)
    panel = panel.merge(dwp, on=["lsoa11", "year"], how="outer")
    panel["dla_pip_source"] = np.where(panel["dla_pip"].notna(), "PLDR", np.where(panel["dwp_dla_pip_count"].notna(), "DWP_August", pd.NA))
    panel["dla_pip"] = panel["dla_pip"].fillna(panel["dwp_dla_pip_count"])
    qof_public = load_qof_practice_lsoa(root)
    panel = panel.merge(qof_public, on=["lsoa11", "year"], how="outer")
    panel["qof_dep_source"] = np.where(panel["qof_dep_pct"].notna(), "PLDR", np.where(panel["qof_public_pct"].notna(), "QOF_practice_LSOA", pd.NA))
    panel["qof_dep_pct"] = panel["qof_dep_pct"].fillna(panel["qof_public_pct"])
    # The hospital component is intentionally absent; do not fill it with EMAS.
    panel["hospital_component"] = np.nan
    for column in ["antidep_rate", "qof_dep_pct", "dla_pip"]:
        panel[f"z_{column}"] = zscore_by_year(panel, column)
    panel["three_component_mean_z"] = panel[["z_antidep_rate", "z_qof_dep_pct", "z_dla_pip"]].mean(axis=1, skipna=False)
    published = load_published_samhi(root / "datasets")
    panel = panel.merge(published, on=["lsoa11", "year"], how="left")
    return panel.sort_values(["year", "lsoa11"]).reset_index(drop=True)


def write_outputs(panel: pd.DataFrame, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    panel.to_csv(output_dir / "three_component_panel.csv", index=False)
    component_cols = ["antidep_rate", "qof_dep_pct", "qof_public_pct", "dla_pip", "hospital_component"]
    coverage = []
    for year, group in panel.groupby("year"):
        for column in component_cols:
            coverage.append({"year": int(year), "component": column, "n": int(group[column].notna().sum()), "coverage_pct": float(group[column].notna().mean() * 100)})
    pd.DataFrame(coverage).to_csv(output_dir / "component_coverage.csv", index=False)
    correlations = []
    for year, group in panel.dropna(subset=["samhi_index"]).groupby("year"):
        for column in ["antidep_rate", "qof_dep_pct", "dla_pip", "three_component_mean_z"]:
            valid = group[[column, "samhi_index"]].dropna()
            correlations.append({"year": int(year), "component": column, "n": len(valid), "pearson_r": valid[column].corr(valid["samhi_index"]) if len(valid) > 1 else np.nan, "spearman_r": valid[column].corr(valid["samhi_index"], method="spearman") if len(valid) > 1 else np.nan})
    pd.DataFrame(correlations).to_csv(output_dir / "component_vs_published_samhi_correlations.csv", index=False)
    LOGGER.info("Wrote component audit outputs to %s", output_dir)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the three observable SAMHI components")
    parser.add_argument("--output-dir", default="scripts/machine_learning/samhi/results/component_reconstruction")
    args = parser.parse_args()
    panel = build_panel(project_root())
    if panel.empty:
        raise SystemExit("No component rows were created")
    write_outputs(panel, project_root() / args.output_dir)


if __name__ == "__main__":
    main()
